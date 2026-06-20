from datetime import date

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_PATH = r"D:\proyecto\CV\Carta_Presentacion_Alfonso_Rubio_Web.docx"


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_bottom_border(paragraph, color="2E74B5", size="8", space="1"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)

    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)

    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)


def configure_page(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def create_style(document, name, font_name, font_size, color, bold=False, italic=False):
    style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = font_name
    style._element.rPr.rFonts.set(qn("w:ascii"), font_name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)
    style.font.size = Pt(font_size)
    style.font.color.rgb = color
    style.font.bold = bold
    style.font.italic = italic
    return style


def set_paragraph_spacing(paragraph, before=0, after=0, line_spacing=1.1):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line_spacing


def spanish_date_today():
    months = {
        1: "enero",
        2: "febrero",
        3: "marzo",
        4: "abril",
        5: "mayo",
        6: "junio",
        7: "julio",
        8: "agosto",
        9: "septiembre",
        10: "octubre",
        11: "noviembre",
        12: "diciembre",
    }
    today = date.today()
    return f"{today.day} de {months[today.month]} de {today.year}"


def add_paragraph(document, text, style_name, alignment=None, before=0, after=0, line_spacing=1.1):
    paragraph = document.add_paragraph(style=style_name)
    if alignment is not None:
        paragraph.alignment = alignment
    set_paragraph_spacing(paragraph, before=before, after=after, line_spacing=line_spacing)
    run = paragraph.add_run(text)
    style = document.styles[style_name]
    set_run_font(
        run,
        name=style.font.name,
        size=style.font.size.pt if style.font.size else None,
        color=style.font.color.rgb,
        bold=style.font.bold,
        italic=style.font.italic,
    )
    return paragraph


def build_document():
    document = Document()
    section = document.sections[0]
    configure_page(section)

    navy = RGBColor(25, 54, 93)
    accent = RGBColor(46, 116, 181)
    muted = RGBColor(96, 96, 96)
    black = RGBColor(0, 0, 0)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)

    create_style(document, "LetterTitle", "Calibri", 22, navy, bold=True)
    create_style(document, "LetterSubtitle", "Calibri", 11.5, muted, bold=False)
    create_style(document, "LetterMeta", "Calibri", 10.5, muted, bold=False)
    create_style(document, "LetterBody", "Calibri", 11, black, bold=False)
    create_style(document, "LetterSubject", "Calibri", 11, navy, bold=True)
    create_style(document, "LetterClosing", "Calibri", 11, black, bold=False)
    create_style(document, "LetterSignature", "Calibri", 11, navy, bold=True)

    add_paragraph(
        document,
        "Carta de presentación",
        "LetterTitle",
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        after=2,
        line_spacing=1.0,
    )
    add_paragraph(
        document,
        "Candidatura para puestos de creación y mantenimiento de páginas web",
        "LetterSubtitle",
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        after=6,
        line_spacing=1.0,
    )
    add_paragraph(
        document,
        "Alfonso Rubio Rioseras | Valladolid | arioseras@hotmail.com | 676 589 037",
        "LetterMeta",
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        after=10,
        line_spacing=1.0,
    )

    rule = document.add_paragraph()
    set_paragraph_spacing(rule, after=14, line_spacing=1.0)
    set_paragraph_bottom_border(rule, color="2E74B5")

    add_paragraph(
        document,
        spanish_date_today(),
        "LetterBody",
        alignment=WD_ALIGN_PARAGRAPH.RIGHT,
        after=12,
    )

    recipient_lines = [
        "[Nombre de la empresa]",
        "At. [Persona responsable de selección]",
        "[Ciudad]",
    ]
    for line in recipient_lines:
        add_paragraph(document, line, "LetterBody", after=0)

    add_paragraph(document, "", "LetterBody", after=6)
    add_paragraph(
        document,
        "Asunto: Candidatura para puestos relacionados con la creación y mantenimiento de páginas web",
        "LetterSubject",
        after=12,
    )
    add_paragraph(
        document,
        "Estimado/a responsable de selección:",
        "LetterBody",
        after=10,
    )

    paragraphs = [
        "Me dirijo a usted para presentar mi candidatura a puestos relacionados con la creación, actualización y mantenimiento de páginas web. Actualmente estoy reforzando mi perfil profesional con formación específica en Confección y Publicación de Páginas Web, y deseo orientar mi trayectoria hacia entornos donde pueda aportar tanto capacidad técnica como compromiso diario con la calidad y el buen funcionamiento de los sitios web.",
        "Cuento con formación como Técnico Superior en Administración de Sistemas Informáticos y con experiencia profesional como operador informático en la Universidad de Valladolid y la Universidad de Burgos. Esta trayectoria me ha permitido trabajar con rigor, resolver incidencias, adaptarme a procedimientos técnicos y mantener una actitud constante de servicio, organización y aprendizaje.",
        "A nivel web, dispongo de base en HTML, CSS y JavaScript, así como conocimientos de herramientas de diseño y maquetación como Photoshop, Illustrator, InDesign y CorelDraw. Considero que esta combinación resulta especialmente útil para colaborar en tareas de edición y actualización de contenidos, mantenimiento de páginas, revisión visual, apoyo a la maquetación y seguimiento del correcto funcionamiento de los recursos web.",
        "Me gustaría formar parte de un equipo en el que seguir creciendo profesionalmente y contribuir con responsabilidad, atención al detalle y una actitud práctica. Quedo a su disposición para ampliar cualquier información en una entrevista personal y agradeceré la oportunidad de valorar mi candidatura.",
    ]

    for text in paragraphs:
        paragraph = add_paragraph(document, text, "LetterBody", after=10, line_spacing=1.1)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    add_paragraph(document, "Atentamente,", "LetterClosing", after=18)
    add_paragraph(document, "Alfonso Rubio Rioseras", "LetterSignature", after=0)

    document.save(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
